#!/usr/bin/env python3
"""Clean a messy CSV/XLSX: trim, normalize headers, drop empty rows/cols, report duplicates."""
import argparse, csv, re, sys


def norm_header(h, i):
    h = re.sub(r"[^0-9a-z]+", "_", str(h or "").strip().lower()).strip("_")
    return h or f"col_{i + 1}"


def read_rows(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def write_rows(path, rows):
    if path.lower().endswith(".xlsx"):
        from openpyxl import Workbook
        wb = Workbook()
        for r in rows:
            wb.active.append(r)
        wb.save(path)
    else:
        with open(path, "w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)


def clean(rows):
    rows = [[re.sub(r"\s+", " ", c).strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return [], 0, []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [i for i in range(width) if any(r[i] for r in rows)]
    rows = [[r[i] for i in keep] for r in rows]
    header = [norm_header(h, i) for i, h in enumerate(rows[0])]
    seen, dups, body = set(), [], []
    for n, r in enumerate(rows[1:], start=2):
        key = tuple(c.lower() for c in r)
        if key in seen:
            dups.append(n)
        seen.add(key)
        body.append(r)
    return [header] + body, width - len(keep), dups


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("input")
    p.add_argument("output")
    p.add_argument("--drop-duplicates", action="store_true", help="remove repeated rows (default: only report)")
    a = p.parse_args()
    rows = read_rows(a.input)
    out, dropped_cols, dups = clean(rows)
    if a.drop_duplicates and dups:
        skip = set(dups)
        out = [out[0]] + [r for n, r in enumerate(out[1:], start=2) if n not in skip]
    write_rows(a.output, out)
    print(f"rows in: {len(rows)}  rows out: {len(out)}  empty cols dropped: {dropped_cols}")
    print(f"duplicate data rows (line numbers after cleaning): {dups or 'none'}")


if __name__ == "__main__":
    sys.exit(main())
